import os, sys
sys.path.insert(0, os.path.dirname(__file__))
from tplkit import Book, E
from pptx import Presentation
from pptx.util import Pt

BASE = r"C:/Users/Rachel/Documents/SEAZONE/BOOK PREMISSAS FACHADA EXTERNA"
TPL = BASE + "/TEMPLATE/TEMPLATE v3 - CAMPECHE SPOT - Book de Premissas (rev13-10-2026).pptx"
C = BASE + "/CANASVIEIRAS SPOT"
OLD = C + "/CANASVIEIRAS SPOT - Book de Premissas de Fachada e Interiores.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else C + "/CANASVIEIRAS SPOT - Book de Premissas_MOODBOARD_v3.pptx"
SCR = os.environ.get("BOOK_CACHE", os.path.join(os.environ.get("TEMP", "."), "book_cache"))

R = C + "/01 RENDERS LANCAMENTO/"
F01, F02, F03, F04, F05, F07, F08 = (R + "01 (1).jpg", R + "02 (2).jpg", R + "03 (2).jpg", R + "04 (2).jpg",
                                     R + "05 (2).jpg", R + "07.jpg", R + "08.jpg")
DRONE = R + "DRONE (1).jpg"
RC = C + "/04 INTERIORES/RECEPÇÃO/"
RT = C + "/04 INTERIORES/ROOFTOP/"
r = lambda n: RC + f"IMAGEM 0{n}.png"
t = lambda n: RT + f"IMAGEM {n:02d}.png"
RC1, RC2, PLANTA, TC1 = RC + "CORREDOR 01.png", RC + "CORREDOR 02.png", RC + "planta recepção.png", RT + "CORREDOR 01.png"

V, D, PD = "A VALIDAR", "DEFINIDO", "PENDENTE"
LEG = "\u2713 definido  \u00b7  ! a validar  \u00b7  ? pendente"

# ---------- marcadores do book antigo (posição normalizada na imagem) ----------
def old_markers(slide_no):
    p = Presentation(OLD)
    sl = p.slides[slide_no - 1]
    pic = [s for s in sl.shapes if s.shape_type == 13 and s.name == "Picture 10"][0]
    cl, cr, ct, cb = pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom
    out = {}
    for s in sl.shapes:
        if s.has_text_frame and s.text_frame.text.strip().isdigit() and s.width < 0.26 * E and \
           pic.left <= s.left <= pic.left + pic.width and pic.top <= s.top <= pic.top + pic.height:
            cx = s.left + s.width / 2; cy = s.top + s.height / 2
            u = cl + (cx - pic.left) / pic.width * (1 - cl - cr)
            v = ct + (cy - pic.top) / pic.height * (1 - ct - cb)
            out[int(s.text_frame.text)] = (u, v)
    return [out[k] for k in sorted(out)]


def place_markers(S, pic_id, win, marker_ids, pts):
    for mid, (u, v) in zip(marker_ids, pts):
        x, y, ok = S.map_pt(pic_id, win, u, v)
        w = S.pos(mid)[2]
        S.geom(mid, x - w / 2, y - w / 2)
        if not ok:
            print("  marcador fora do quadro:", mid)
    for mid in marker_ids[len(pts):]:
        S.delete(mid)


def label_w(S, sid, txt):
    S.T(sid, txt)
    S.geom(sid, w=0.30 + 0.078 * len(txt))


def card(S, ids, code, title, onde, spec, obra, img, focus, hfrac):
    pic, lab, txt = ids
    S.P(pic, img, focus, hfrac)
    S.T(lab, code)
    S.T(txt, title, "ONDE", onde, "ESPECIFICAÇÃO", spec, "ATENÇÃO NA OBRA", obra)


def rows_with_pills(S, rows, data, pill_ids, full_w=5.98, short_w=4.98):
    """rows=[(label_id,text_id)], data=[(label,text,status|None)]"""
    for pid in pill_ids:
        S.delete(pid)
    for (lid, tid), (lab, txt, st) in zip(rows, data):
        S.T(lid, lab); S.T(tid, txt)
        S.geom(tid, w=short_w if st else full_w)
        if st:
            x, y, w, h = S.pos(tid)
            pw = {"A VALIDAR": 0.97, "DEFINIDO": 0.90, "PENDENTE": 0.90, "A COTAR": 0.82}[st]
            S.pill(st, 12.71 - pw, y + h / 2 - 0.12)


def checklist(S, data):
    rows = [(775, 776, 777), (778, 779, 780), (782, 783, 784), (785, 786, 787), (789, 790, 791), (792, 793, 794)]
    for (box, tid, pid), (title, body, st) in zip(rows, data):
        S.T(tid, title, body)
        y = S.pos(pid)[1]
        S.delete(pid)
        pw = {"A VALIDAR": 0.97, "DEFINIDO": 0.90, "PENDENTE": 0.90, "A COTAR": 0.82}[st]
        S.pill(st, 12.71 - pw, y)


def vista_items(S, pairs, data):
    """pairs=[(num_id,text_id,icon_id)]"""
    for (nid, tid, iid), (title, body, st) in zip(pairs, data):
        S.T(tid, title, body)
        S.status_icon(iid, st)
    for (nid, tid, iid) in pairs[len(data):]:
        S.delete(nid, tid, iid)


B = Book(TPL, SCR + "/cache_canas")

# ================= 1 · CAPA =================
S = B.dup(1)
S.P(16, F05, (0.47, 0.5))
S.P(17, t(2), (0.5, 0.5))
S.P(18, r(2), (0.42, 0.5))
S.P(19, F04, (0.42, 0.62))
S.T(22, "Canasvieiras Spot")
S.sh(22).text_frame.paragraphs[0].runs[0].font.size = Pt(36)
S.T(25, "Canasvieiras · Florianópolis/SC", "DOCUMENTO TÉCNICO INTERNO")

# ================= 2 · CONTEXTUALIZAÇÃO =================
S = B.dup(2)
S.T(36, "Consolidar as premissas de fachada e de interiores das áreas comuns do Canasvieiras Spot a partir da "
        "leitura cruzada entre as imagens de lançamento, o projeto executivo de arquitetura e o projeto de "
        "interiores da Seazone: o que a imagem vendeu, o que o executivo especifica, onde divergem e o que "
        "validar antes da compra e da execução.")
S.P(37, F07, (0.52, 0.5), 0.55)
S.P(40, F01, (0.43, 0.252), 0.05)
S.P(43, F01, (0.25, 0.86), 0.16)
qs = ["O que as imagens de lançamento mostram na fachada?",
      "Qual material e cor o executivo define para cada volume?",
      "Onde o executivo diverge da imagem vendida?",
      "Quais cores e texturas exigem amostra aprovada?",
      "Qual é o conceito de interiores das áreas comuns que deve ser seguido?",
      "O que não pode ser alterado sem validação?",
      "Quais pontos ainda precisam ser validados?"]
for sid, q in zip([49, 51, 53, 55, 57, 59, 61], qs):
    S.T(sid, q)
S.T(62, "Regra do material: nenhuma informação foi inventada. Fontes (Canasvieiras Spot / 02 - Projetos): imagens "
        "de lançamento, executivo de arquitetura (elevações, detalhes, esquadrias e brises), Memorial Descritivo, "
        "projeto de interiores Seazone (Memorial de Texturas e Memorial de Interiores) e 00 - Premissas de Projeto "
        "(Briefing e LOG de considerações). Foco: imagens externas, texturas e interiores das áreas comuns. "
        "Interiores das unidades, índices urbanísticos, áreas e custos não fazem parte deste Book.")

# ================= 3 · COMO USAR =================
S = B.dup(3)
S.P(75, F01, (0.43, 0.252), 0.05)
S.T(77, "REV-08")
S.T(78, "Código", "Código do executivo/memorial (ex.: REV-08, BRI-01, T07) para achar o item no projeto.")
S.P(87, F01, (0.5, 0.45), None, 0.75)
S.P(90, t(3), (0.5, 0.5))
S.P(93, t(5), (0.3, 0.45))
S.T(89, "Fachada", "imagens de lançamento · paleta · leitura geral · texturas T01–T12 · vistas de esquina, "
                   "posterior e térreo · pendências")
S.T(92, "Interiores das áreas comuns", "conceito × situação · recepção e corredores · rooftop interno e externo · "
                                       "pendências")
S.T(95, "Materiais de interiores das áreas comuns", "texturas T13–T24 para compra e aprovação de amostras")

# ================= 4 · DIVISOR 01 =================
S = B.dup(4)
S.P(101, DRONE, (0.45, 0.5))
S.T(105, "Imagens de lançamento  ·  paleta  ·  leitura geral  ·  texturas  ·  vistas com marcação e "
         "especificação  ·  pendências")

# ================= 5 · MOODBOARD FACHADA =================
S = B.dup(5)
S.P(118, F01); label_w(S, 119, "01 · FRONTAL")
S.P(120, F07); label_w(S, 121, "07 · ESQUINA")
S.P(122, F02, (0.45, 0.5)); label_w(S, 123, "02")
S.P(124, F05, (0.47, 0.5)); label_w(S, 125, "05")
S.P(126, F04, (0.45, 0.5)); label_w(S, 127, "04 · TÉRREO")
S.P(128, t(3)); label_w(S, 129, "ROOFTOP")
for sid, (a, b) in zip([132, 134, 136, 138, 140], [
        ("Volume central amadeirado", "porcelanato Boreal Natural (REV-08) com brises verticais de correr nas janelas"),
        ("Volumes laterais", "cinza escuro (REV-06) e branco (REV-07) com floreiras nas bordas de laje"),
        ("Brises e elementos vazados", "brises pretos e Oak (BRI-01), painéis em grelha (EV3/EV4) e cobogó sobre as lojas (EV1)"),
        ("Base e acesso", "lojas com vitrines, brises Oak no mezanino e totem iluminado CANASVIEIRAS SPOT no muro"),
        ("Coroamento e rooftop", "guarda-corpo de vidro, pergolado preto, piscina e letreiro SEAZONE no reservatório")]):
    S.T(sid, a, b)

# ================= 6 · PALETA =================
S = B.dup(6)
pal = [(151, 152, 153, "REV-08", "Porcelanato Boreal Natural", "volume central e montantes", F01, (0.43, 0.252), 0.05),
       (154, 155, 156, "REV-06", "Cinza escuro Rock'n Roll", "volumes laterais e bases", F02, (0.60, 0.34), 0.07),
       (157, 158, 159, "REV-07", "Branco Puro", "volumes claros e reservatório", F02, (0.67, 0.62), 0.15),
       (160, 161, 162, "REV-05", "Cinza ref. Nanquim", "faixas de laje e marquise", F05, (0.39, 0.66), 0.07),
       (163, 164, 165, "BRI-01", "Brises preto e Oak", "janelas, térreo e mezanino", F05, (0.335, 0.75), 0.07),
       (166, 167, 168, "EV1–EV4", "Elemento vazado 40×40", "lojas e painéis laterais", F01, (0.72, 0.70), 0.09)]
for pic, lab, txt, code, a, b, img, foc, hf in pal:
    S.P(pic, img, foc, hf); S.T(lab, code); S.T(txt, a, b)
S.T(170, "COMO LER", "Fachada = volume central amadeirado (REV-08) + volumes laterais cinza escuro (REV-06) e branco "
                     "(REV-07) + faixas cinza (REV-05) + brises pretos e Oak + elementos vazados + esquadria preta + "
                     "floreiras. Detalhe de cada material nas páginas de texturas T01–T12.")

# ================= 7 · LEITURA GERAL =================
S = B.dup(7)
win = S.P(181, F01)
S.T(182, "LANÇAMENTO · 01 (FRONTAL)")
place_markers(S, 181, win, list(range(183, 193)), old_markers(4))
S.T(193, "A fachada é lida em quatro camadas: volume central em porcelanato amadeirado (REV-08) com brises verticais "
         "de correr; volumes laterais em cinza escuro (REV-06) e branco (REV-07) com floreiras e painéis em elemento "
         "vazado; base do térreo e mezanino com brises Oak, elemento vazado sobre as lojas e totem de acesso; "
         "coroamento com guarda-corpo de vidro, pergolado e letreiro SEAZONE.")
leg = [("Letreiro SEAZONE no coroamento", "letras-caixa iluminadas · não detalhado no executivo"),
       ("Pergolado da cobertura", "AC-02 pergolado em alumínio"),
       ("Guarda-corpo dos terraços", "GC-01 alumínio / vidro"),
       ("Volume central amadeirado", "REV-08 porcelanato Boreal Natural 20×120"),
       ("Brises verticais das janelas", "BRI-01 brise de correr 75×25 · preto e Oak"),
       ("Ripado horizontal nas janelas", "só no render · não consta nas elevações"),
       ("Painéis laterais em grelha", "EV3 elemento vazado de concreto 40×40"),
       ("Floreiras", "AC-01 vaso · floreira Vasoflex"),
       ("Brises amadeirados do mezanino", "Lesco Green Brise · Oak"),
       ("Elemento vazado sobre as lojas", "EV1 concreto 40×40")]
for sid, (a, b) in zip(range(196, 215, 2), leg):
    S.T(sid, a, b)

# ================= 8-10 · TEXTURAS DA FACHADA =================
CARDS = [(227, 228, 229), (231, 232, 233), (235, 236, 237), (239, 240, 241)]
NOTE = ("Amostras extraídas das imagens de lançamento: referência visual, não substitui amostra física. "
        "Especificação: executivo de arquitetura, Memorial Descritivo e Briefing.")
tex = [
    ("Texturas da fachada · cores e revestimento (1/3)", [
        ("T01", "REV-06 · Cinza escuro Rock'n Roll", "Volumes escuros laterais, bases e molduras nas quatro elevações.",
         "Pintura acrílica sobre textura/massa acrílica, ref. cor Rock'n Roll (Suvinil).",
         "Aprovar amostra em painel na própria fachada antes da compra; mesma cor e lote nas quatro faces.",
         F02, (0.60, 0.34), 0.05),
        ("T02", "REV-05 · Cinza ref. Nanquim", "Faixas de laje, marquise do térreo e fundo do elemento vazado das lojas.",
         "Pintura acrílica descrita como \"cinza claro\", ref. cor Nanquim (Suvinil).",
         "Nome da cor e descrição não conferem entre si: validar amostra contra o render. A VALIDAR.",
         F05, (0.39, 0.695), 0.035),
        ("T03", "REV-07 · Branco Puro", "Volumes laterais claros, platibanda da cobertura e caixa do reservatório.",
         "Pintura acrílica, ref. Branco Puro (Suvinil).",
         "Reservatório com acabamento externo idem fachada (memorial 30.4.3).",
         F02, (0.67, 0.62), 0.08),
        ("T04", "REV-08 · Porcelanato amadeirado", "Volume central das fachadas frontal e posterior e montantes entre janelas.",
         "Porcelanato retificado 20×120, Boreal Natural (Portobello) · memorial 10.1 e 30.3.5.",
         "O render mostra tom mais quente e claro: aprovar peça física contra a imagem e definir paginação. A VALIDAR.",
         F01, (0.43, 0.252), 0.05)]),
    ("Texturas da fachada · brises e elementos vazados (2/3)", [
        ("T05", "BRI-01 · Brise amadeirado Oak", "Térreo, mezanino e parte dos pavimentos 2 a 4.",
         "Brise de correr em alumínio, ref. Lesco Green Brise 75×25 mm, cor Oak (prancha 28).",
         "Briefing (gerais) marca brise em PVC e desmarca metálico: confirmar material antes de comprar. A VALIDAR.",
         F05, (0.335, 0.75), 0.05),
        ("T06", "BRI-01 · Brise alumínio preto", "Janelas dos pavimentos 2, 3 e 4 e do mezanino.",
         "Brise de correr em alumínio preto 75×25 mm · trilho BRI-02 preto · BRI-03 veneziana preta nas condensadoras.",
         "Tipos mudam por pavimento: conferir o quadro da prancha 28 antes de fabricar. Material × briefing (PVC): A VALIDAR.",
         F01, (0.365, 0.36), 0.05),
        ("T07", "EV1 · Elemento vazado das lojas", "Fachada sobre as lojas e EV2 no térreo.",
         "Elemento vazado de concreto 40×40, peças fixas (quadro de esquadrias, prancha 27).",
         "O render mostra borda inferior \"pixelada\" (peças omitidas); o executivo mostra borda reta. A VALIDAR.",
         F01, (0.72, 0.70), 0.07),
        ("T08", "EV3/EV4 · Painéis em grelha", "Faixas verticais ao lado do volume central.",
         "Elemento vazado de concreto 40×40; cor da peça não especificada.",
         "No render as peças são pretas: definir se recebem REV-06 ou ficam aparentes. A VALIDAR.",
         F01, (0.32, 0.47), 0.06)]),
    ("Texturas da fachada · esquadrias e cobertura (3/3)", [
        ("T09", "Esquadrias de alumínio", "Janelas e porta-janelas de todas as fachadas, lojas e áreas comuns.",
         "Alumínio anodizado preto (padrão Alcoa) · vidro laminado transparente (memorial 16.1).",
         "Peitoril e pingadeira em granito preto São Gabriel com frisos (memorial 20); quadro nas pranchas 26 e 27.",
         F01, (0.255, 0.36), 0.06),
        ("T10", "GC-01/02/03 · Guarda-corpos", "Terraços da cobertura e borda da piscina.",
         "Alumínio anodizado preto + vidro (executivo).",
         "Briefing: vidro incolor com perfil \"U\" inox, sem perfis de alumínio × executivo: alumínio + vidro. A VALIDAR.",
         F01, (0.225, 0.225), 0.045),
        ("T11", "AC-01 · Floreiras de fachada", "Bordas de laje dos volumes laterais nos pavimentos 2 a 4.",
         "AC-01 vaso para plantas · floreira modelo estrutural Vasoflex (Ecotelhado).",
         "Fixação mecânica com buchas e vasos independentes; paisagismo por empresa especializada.",
         F07, (0.445, 0.36), 0.06),
        ("T12", "AC-02 · Pergolado da cobertura", "Áreas de integração e terraços do último pavimento.",
         "AC-02 pergolado em alumínio (elevações) · memorial 30.3.5: \"estrutura metálica com madeira\".",
         "LOG Ind. 06: reforço na estrutura para o balanço. Alumínio × metálico com madeira: A VALIDAR.",
         F01, (0.36, 0.16), 0.04)])]
for title, cards in tex:
    S = B.dup(8)
    S.T(221, title); S.T(225, NOTE)
    for ids, (code, ttl, onde, spec, obra, img, foc, hf) in zip(CARDS, cards):
        card(S, ids, code, ttl, onde, spec, obra, img, foc, hf)

# ================= 11 · VISTA ESQUINA (tpl 13) =================
V13 = [(430, 431, 432), (433, 434, 435), (436, 437, 438), (439, 440, 441), (442, 443, 444),
       (445, 446, 447), (448, 449, 450), (451, 452, 453), (454, 455, 456), (457, 458, 459)]
M13 = list(range(416, 426))
S = B.dup(13)
S.T(410, "Esquina: fachada frontal e lateral")
win = S.P(414, F07)
label_w(S, 415, "LANÇAMENTO · 07 (ESQUINA)")
place_markers(S, 414, win, M13, old_markers(7))
S.T(427, "LEITURA DA ESQUINA", "A esquina concentra os três materiais de maior impacto visual: porcelanato amadeirado "
         "(REV-08), cinza escuro (REV-06) e brises. As floreiras marcam os volumes laterais e o elemento vazado com "
         "borda \"pixelada\" fecha o volume das lojas. O totem iluminado não tem detalhe no executivo.")
S.T(429, LEG)
vista_items(S, V13, [
    ("Letreiro SEAZONE", "Letras-caixa iluminadas no reservatório (render) · sem detalhe no executivo", PD),
    ("Pergolado", "AC-02 em alumínio · perfis pretos no render", V),
    ("Guarda-corpo", "GC-01 alumínio anodizado preto e vidro", D),
    ("Volume central", "REV-08 porcelanato 20×120 Boreal Natural (Portobello)", D),
    ("Brises pretos", "BRI-01 de correr, alumínio preto 75×25 + trilho BRI-02", D),
    ("Floreiras", "AC-01 Vasoflex, fixação mecânica", D),
    ("Painel em grelha", "EV3 elemento vazado de concreto 40×40 · cor a definir", V),
    ("Brises Oak", "BRI-01 Lesco Green Brise Oak · térreo e mezanino", D),
    ("Elemento vazado", "EV1 concreto 40×40 sobre REV-05 · borda a confirmar", V),
    ("Totem de acesso", "\"CANASVIEIRAS SPOT\" em letras iluminadas no muro (render)", V)])

# ================= 12 · VISTA POSTERIOR (tpl 12) =================
S = B.dup(12)
S.T(354, "Fachada posterior e lateral")
win = S.P(358, F05)
label_w(S, 359, "LANÇAMENTO · 05 (POSTERIOR)")
place_markers(S, 358, win, list(range(360, 369)), old_markers(8)[:7])
S.T(370, LEG)
V12 = [(371, 372, 373), (374, 375, 376), (377, 378, 379), (380, 381, 382), (383, 384, 385),
       (386, 387, 388), (389, 390, 391), (392, 393, 394), (395, 396, 397)]
vista_items(S, V12, [
    ("Piscina da cobertura", "Borda com guarda-corpo GC-03 · Barlavento Blanc 14,5×14,5 (REV-09) × briefing: esverdeado 20×20", V),
    ("Volume amadeirado", "REV-08 Boreal Natural (Portobello) · mesma leitura da fachada frontal", D),
    ("Brises das janelas", "BRI-01 preto e Oak · trilho BRI-02 · BRI-03 veneziana preta nas condensadoras", D),
    ("Ripado horizontal", "No peitoril das janelas (render) · não consta nas elevações", V),
    ("Brises Oak do térreo", "Brise 3 e Brise 4 Oak no pavimento 1", D),
    ("Volume escuro lateral", "REV-06 cinza escuro, ref. Rock'n Roll (Suvinil)", D),
    ("Montantes brancos", "REV-07 Branco Puro (Suvinil)", D)])
S.P(398, F01, (0.43, 0.252), 0.05)
S.P(399, F02, (0.60, 0.34), 0.05)
S.P(400, F02, (0.67, 0.62), 0.08)
S.T(401, "REV-08 · REV-06 · REV-07")
S.T(403, "NAS QUATRO ELEVAÇÕES", "As fachadas laterais e posterior repetem a paleta da frontal: REV-05, REV-06, "
         "REV-07 e REV-08 + brises. Referências: elevações (pranchas 19 a 22) e detalhes de elevação (pranchas 23 a 25).")

# ================= 13 · VISTA TÉRREO (tpl 13) =================
S = B.dup(13)
S.T(410, "Térreo: acesso, lojas e totem")
win = S.P(414, F04)
label_w(S, 415, "LANÇAMENTO · 04 (TÉRREO E MEZANINO)")
place_markers(S, 414, win, M13, old_markers(9))
S.T(427, "ATENÇÃO · IDENTIDADE DO ACESSO", "O totem iluminado CANASVIEIRAS SPOT no muro, o cobogó das lojas e os "
         "brises Oak do mezanino são a \"cara\" do térreo no lançamento. O executivo traz o nome no muro sem detalhe "
         "de letra, iluminação e acabamento: detalhar antes da execução.")
S.T(429, LEG)
vista_items(S, V13, [
    ("Totem CANASVIEIRAS SPOT", "Letras iluminadas no muro · sem detalhe no executivo", V),
    ("Rampa de acesso", "Guarda-corpo e corrimão em perfil preto", D),
    ("Paisagismo do passeio", "Vegetação alta em frente ao térreo", D),
    ("Lojas 01 e 02", "Vitrines em vidro · esquadria de alumínio preto", D),
    ("Elemento vazado EV1", "Concreto 40×40 · borda \"pixelada\" no render", V),
    ("Brises Oak do mezanino", "Brise 2 Oak (Lesco Green Brise)", D),
    ("Painel em grelha EV3", "Elemento vazado 40×40 · cor a definir", V),
    ("Floreiras AC-01", "Vaso · floreira Vasoflex", D),
    ("Faixa de laje / marquise", "REV-05 pintura cinza (ref. Nanquim)", V),
    ("Esquadrias do pavimento 1", "Alumínio anodizado preto · vidro", D)])

# ================= 14 · PENDÊNCIAS DA FACHADA (tpl 24) =================
S = B.dup(24)
S.T(767, "01 · FACHADA · PENDÊNCIAS")
S.T(768, "Fachada: pendências e pontos a confirmar")
S.P(772, F02, (0.45, 0.5))
label_w(S, 773, "LANÇAMENTO · 02")
checklist(S, [
    ("Letreiro SEAZONE no coroamento", "Detalhe, material, iluminação e fixação no reservatório (memorial 28.2 sem detalhe).", PD),
    ("Ripado horizontal das janelas", "Confirmar se o ripado do render entra na obra ou se prevalece o brise de piso a teto do executivo.", V),
    ("Elemento vazado EV1–EV4", "Cor das peças (render preto × concreto aparente ou REV-06) e borda \"pixelada\" do EV1 sobre as lojas.", V),
    ("Brises e guarda-corpos", "Briefing: brise em PVC e vidro com perfil \"U\" inox × executivo: alumínio Lesco e alumínio + vidro.", V),
    ("Cores e porcelanato", "REV-05 \"cinza claro\" × ref. Nanquim; tom do Boreal Natural (REV-08) contra o render: aprovar amostras.", V),
    ("Pergolado, totem e soleiras", "Pergolado alumínio × metálico com madeira; detalhe do totem; soleira granito cinza (pranchas) × porcelanato (memorial 20.1).", V)])

# ================= 15 · DIVISOR 02 =================
S = B.dup(14)
S.P(465, r(1), (0.62, 0.5))
S.T(468, "Interiores das áreas comuns")
S.T(469, "Conceito Seazone × situação  ·  recepção e corredores  ·  rooftop interno  ·  rooftop externo e piscina  ·  pendências")

# ================= 16 · CONCEITO × SITUAÇÃO (tpl 15) =================
S = B.dup(15)
cards15 = [(483, 484, 485, 486), (488, 489, 490, 491), (493, 494, 495, 496), (498, 499, 500, 501), (503, 504, 505, 506)]
data15 = [
    ("Recepção / hall", r(2), (0.45, 0.5),
     "Painel do logo em Velvet off-white com \"Canasvieiras SPOT\" e LED no \"O\", sarrafos e floreiras em MDF Savana, sofá em couro",
     "Paredes: Vela / Camurça Clara (texturas) × briefing \"Outono Gelado\" (Coral)", V),
    ("Corredores", RC1, (0.55, 0.5),
     "Paredes off-white, sarrafos verticais em MDF Savana, arandelas cilindro com haste de madeira, piso claro 90×90",
     "Portas das unidades brancas (briefing e memorial; imagem ilustrativa) · rodapé branco × preto (texturas)", V),
    ("Rooftop interno", t(5), (0.3, 0.5),
     "Gourmet com pedra moledo, ripado de madeira, armários cinza, bancada Pitaya e pendente de rattan",
     "Marcenaria do rooftop \"a definir\" no memorial de interiores × Savana e Cinza Sagrado (texturas)", PD),
    ("Rooftop externo", t(3), (0.5, 0.5),
     "Piscina com deck amadeirado, painel do logo em pedra, balanço e pergolado preto",
     "Revestimento da piscina e fixação do balanço divergem entre memorial, briefing e LOG", V),
    ("Projeto de interiores", t(10), (0.5, 0.5),
     "Imagens Seazone + Memorial de Texturas + Memorial de Interiores (10 - Projeto de Interiores)",
     "Recepção entregue mobiliada (memorial 31) · marcenaria, bancadas e luminárias pela Decor (memorial 32.1)", D)]
for (pic, ttl, body, pill), (name, img, foc, conc, sit, st) in zip(cards15, data15):
    S.P(pic, img, foc)
    S.T(ttl, name)
    S.T(body, "CONCEITO SEAZONE (SEGUIR)", conc, "SITUAÇÃO REGISTRADA", sit)
    x, y, w, h = S.pos(pill)
    S.delete(pill)
    S.pill(st, x, y)
S.T(507, "Imagens Seazone (10 - Projeto de Interiores / 02 - Imagens) × Memorial Descritivo, Memorial de Texturas, "
         "Memorial de Interiores, Briefing e LOG de considerações (00 - Premissas de Projeto).")
S.T(509, "ATENÇÃO · REGRA PARA A OBRA", "Vale o conceito Seazone das áreas comuns. Onde a imagem mostra cimento "
         "queimado, aplicar textura Velvet (memorial de texturas). Qualquer troca de material ou mobiliário precisa de "
         "validação registrada.")

# ================= 17 · RECEPÇÃO MOODBOARD (tpl 25) =================
S = B.dup(25)
S.T(800, "02 · INTERIORES · RECEPÇÃO E CORREDORES · MOODBOARD")
S.T(801, "Recepção e corredores: Velvet, Savana e couro")
for pic, img, foc in zip(range(805, 814), [r(1), r(2), r(3), r(4), r(5), RC1, RC2, PLANTA, TC1],
                         [(0.55, .5), (0.45, .5), (0.4, .5), (0.5, .5), (0.5, .5), (0.55, .5), (0.5, .5), (0.5, .5), (0.5, .5)]):
    S.P(pic, img, foc)
S.T(816, "CONCEITO", "Recepção com painel do logo em textura Velvet off-white e letreiro \"Canasvieiras SPOT\" com LED "
         "no \"O\", sarrafos verticais e floreiras em MDF Savana e muita vegetação.",
         "CORREDORES", "Paredes off-white acetinadas com sarrafos em MDF Savana, arandelas cilindro com haste de "
         "madeira e piso claro 90×90.",
         "MOBILIÁRIO", "Sofá em couro, puffs Aurora e Saturno cinza, puff Seat Garden e tapete Debrum Chevron cinza "
         "(memorial de interiores).")
S.T(815, "REGRA DO MEMORIAL DE TEXTURAS", "Onde a imagem mostra cimento queimado, aplicar textura Velvet. Portas das "
         "unidades nos corredores: brancas.")
for pic, lab, name, img, foc, hf in [(817, 818, "Velvet off-white", r(2), (0.30, 0.12), 0.15),
                                      (819, 820, "Superquadra", r(2), (0.90, 0.88), 0.12),
                                      (821, 822, "MDF Savana", RC1, (0.81, 0.45), 0.12),
                                      (823, 824, "Couro", r(2), (0.50, 0.62), 0.10)]:
    S.P(pic, img, foc, hf); S.T(lab, name)

# ================= 18 · FICHA RECEPÇÃO (tpl 27) =================
S = B.dup(27)
S.T(848, "02 · INTERIORES · RECEPÇÃO E CORREDORES")
S.T(849, "Ficha: recepção, hall e corredores")
S.P(853, r(2), (0.45, 0.5)); label_w(S, 854, "CONCEITO · RECEPÇÃO")
S.P(855, RC1, (0.55, 0.5)); label_w(S, 856, "CONCEITO · CORREDOR")
S.T(857, "Recepção entregue mobiliada e equipada (memorial 31). Imagens em 10 - Projeto de Interiores / 02 - Imagens / RECEPÇÃO.")
rows_with_pills(S, [(859, 860), (861, 862), (865, 866), (867, 868), (870, 871), (873, 874), (876, 877), (878, 879), (881, 882)], [
    ("Painel do logo", "Textura Velvet off-white (Suvinil Efeito Veludo / Coral Design) + \"Canasvieiras SPOT\" preto com LED no \"O\"", None),
    ("Paredes", "Off-white acetinado · Suvinil \"Vela\" / Coral \"Camurça Clara\" (texturas) × briefing corredores: \"Outono Gelado\" (Coral)", V),
    ("Ripado e floreiras", "Sarrafos verticais em MDF Savana com verniz fosco · floreiras em MDF Savana", None),
    ("Piso", "Porcelanato Superquadra Concreto Natural 90×90 retificado (Portobello) · lobby e corredores", None),
    ("Teto", "Recepção: gesso acartonado com pé-direito duplo (briefing) · corredores: forro modular mineral (memorial 15.2)", None),
    ("Rodapé", "Briefing e prancha 24: poliestireno branco de sobrepor, quina viva × memorial de texturas: esmalte preto", V),
    ("Mobiliário", "Sofá em couro · puffs Aurora e Saturno cinza · puff Seat Garden · tapete Debrum Chevron cinza", None),
    ("Portas", "Portas das unidades brancas, bandeira até o forro (briefing) = laca PU branca (memorial 17.1)", None),
    ("Iluminação", "Trilho com spots direcionáveis em alumínio preto · arandelas de corredor: cilindro off-white + haste Savana", None)],
    [863, 872])

# ================= 19 · ROOFTOP MOODBOARD (tpl 28) =================
S = B.dup(28)
S.T(888, "02 · INTERIORES · ROOFTOP · MOODBOARD")
S.T(889, "Rooftop: integração, gourmet e piscina")
for pic, n in zip(range(893, 902), [5, 6, 7, 8, 9, 10, 1, 2, 3]):
    S.P(pic, t(n))
S.T(902, "CONCEITO", "Gourmet e copa com pedra moledo, ripados de madeira, armários cinza e pendente de rattan; deck "
         "amadeirado com piscina e painel do logo em pedra.",
         "REGRA DE TEXTURA", "Onde a imagem mostra cimento queimado: textura Velvet cinza concreto (memorial de texturas).",
         "MOBILIÁRIO", "Cadeira Mara bouclé cinza, poltrona em alumínio preto, banqueta Clarice, balanço Acapulco e "
         "espreguiçadeira em aço preto.")
for pic, lab, name, img, foc, hf in [(903, 904, "Pedra moledo", t(5), (0.12, 0.25), 0.15),
                                      (905, 906, "Velvet cinza", t(2), (0.10, 0.30), 0.15),
                                      (907, 908, "Ripado", t(5), (0.40, 0.30), 0.18),
                                      (909, 910, "Deck", t(1), (0.45, 0.88), 0.12)]:
    S.P(pic, img, foc, hf); S.T(lab, name)

# ================= 20 · FICHA ROOFTOP INTERNO (tpl 30) =================
S = B.dup(30)
S.T(938, "02 · INTERIORES · ROOFTOP INTERNO")
S.T(939, "Ficha de premissas: gourmet, copa e mini market")
S.P(943, t(5), (0.3, 0.5)); label_w(S, 944, "CONCEITO SEAZONE · GOURMET")
S.P(945, t(9), (0.5, 0.5)); label_w(S, 946, "MINI MARKET")
S.T(947, "Memorial de interiores lista a marcenaria do rooftop como \"material a definir\": fechar referência antes do pedido.")
rows_with_pills(S, [(949, 950), (952, 953), (955, 956), (957, 958), (960, 961), (962, 963), (965, 966), (967, 968), (970, 971)], [
    ("Painéis cinza", "Textura Velvet cinza concreto médio (Suvinil \"Cimento Queimado\"/\"Concreto\") · aprovar amostra contra o render", None),
    ("Paredes internas", "Pintura fosco aveludado lavável cinza claro · Suvinil \"Cromo\" / Coral \"Cinza Sublime\"", None),
    ("Pedra da copa", "Pedra branca estilo moledo, irregular, junta seca (briefing e memorial de texturas)", None),
    ("Piso interno", "Porcelanato 90×90 natural/acetinado (briefing)", None),
    ("Marcenaria", "Ripados e prateleiras em MDF Savana · armários do gourmet e mini market em MDF Cinza Sagrado × memorial de interiores \"a definir\"", PD),
    ("Bancada da copa", "Granito Pitaya · cuba inox · torneira Gali bica alta (Docol)", None),
    ("Forro", "Forro modular no rooftop e no mini market (briefing)", None),
    ("Mobiliário", "Cadeira Mara bouclé cinza · banqueta Clarice · poltrona em alumínio preto", None),
    ("Iluminação", "Spots de sobrepor pretos · pendente Bulan em rattan", None)],
    [951])

# ================= 21 · FICHA ROOFTOP EXTERNO (tpl 32) =================
S = B.dup(32)
S.T(1007, "02 · INTERIORES · ROOFTOP EXTERNO")
S.T(1008, "Ficha de premissas: rooftop externo e piscina")
S.P(1012, t(3), (0.5, 0.5)); label_w(S, 1013, "CONCEITO · PISCINA")
S.P(1014, t(2), (0.45, 0.5)); label_w(S, 1015, "CONCEITO · BALANÇO E LOGO")
S.T(1016, "Painel do logo em pedra moledo com \"Canasvieiras SPOT\"; floreira e banco em alvenaria com a mesma pedra (LOG Ind. 02).")
rows_with_pills(S, [(1018, 1019), (1021, 1022), (1024, 1025), (1026, 1027), (1029, 1030), (1031, 1032), (1034, 1035), (1037, 1038), (1041, 1042)], [
    ("Piscina", "Barlavento Blanc 14,5×14,5 (memorial 30.3.1, prancha 44) × briefing: revestimento esverdeado 20×20", V),
    ("Borda", "Borda curva (prancha) × borda infinita/atérmica: definir acabamento", V),
    ("Aquecimento", "Solar (briefing e memorial)", None),
    ("Deck", "Porcelanato que imita madeira 120×20 (briefing) sobre piso elevado (prancha 44) · referência não definida", PD),
    ("Pedra", "Pedra branca estilo moledo na parede do balanço e lateral · floreira e banco com a mesma pedra (LOG Ind. 02)", None),
    ("Painel do logo", "\"Canasvieiras SPOT\" preto com LED no \"O\" sobre a pedra · ponto com interruptor no logo (briefing)", None),
    ("Balanço", "Reforço no pergolado (LOG Ind. 06) × instalação direto na laje (briefing)", V),
    ("Guarda-corpo", "GC-03 alumínio anodizado + vidro (executivo) × briefing: vidro com perfil \"U\" inox", V),
    ("Mobiliário", "Balanço Acapulco · espreguiçadeira em aço preto · banqueta Clarice", None)],
    [1020, 1036, 1039])

# ================= 22 · PENDÊNCIAS INTERIORES (tpl 24) =================
S = B.dup(24)
S.T(767, "02 · INTERIORES · PENDÊNCIAS")
S.T(768, "Áreas comuns: pendências e pontos a confirmar")
S.P(772, t(1), (0.78, 0.5))
label_w(S, 773, "ROOFTOP · BALANÇO")
checklist(S, [
    ("Paredes das áreas comuns", "Vela / Camurça Clara (memorial de texturas) × \"Outono Gelado\" (briefing): fechar a cor dos corredores.", V),
    ("Rodapé", "Branco de sobrepor (briefing e prancha 24) × esmalte preto (memorial de texturas).", V),
    ("Revestimento e borda da piscina", "Barlavento Blanc (memorial) × esverdeado 20×20 (briefing) · borda curva × infinita.", V),
    ("Balanço do rooftop", "Reforço no pergolado (LOG Ind. 06) × instalação na laje (briefing).", V),
    ("Deck do rooftop", "Porcelanato que imita madeira 120×20: falta a referência comercial antiderrapante.", PD),
    ("Marcenaria do rooftop e da recepção", "Materiais \"a definir\" no memorial de interiores × Savana e Cinza Sagrado (texturas).", PD)])

# ================= 23 · DIVISOR 03 =================
S = B.dup(33)
S.P(1048, t(5), (0.3, 0.45))
S.T(1052, "Áreas comuns · mapa de materialidade T13–T24: amostras tiradas das imagens Seazone para compra e aprovação")

# ================= 24-26 · MATERIAIS =================
MC = [(1067, 1068, 1069), (1071, 1072, 1073), (1075, 1076, 1077), (1079, 1080, 1081)]
mats = [
    ("Materiais · pisos e paredes (1/3)", [
        ("T13", "Porcelanato Superquadra", "Lobby, recepção, corredores, WC comuns e terraços da cobertura.",
         "Superquadra Concreto Natural 90×90 retificado (Portobello) · memorial 12.4 e 30.",
         "Peças grandes: argamassa também no tardoz; alinhar paginação hall × circulação.", r(2), (0.90, 0.88), 0.12),
        ("T14", "Textura Velvet", "Onde a imagem mostra cimento queimado: painel do logo da recepção e painéis cinza do rooftop.",
         "Suvinil Efeito Veludo / Coral Design Efeito Veludo · off-white (recepção) e cinza concreto (rooftop).",
         "Aprovar amostra contra o render antes de aplicar em área grande.", t(2), (0.10, 0.30), 0.15),
        ("T15", "Pintura das paredes", "Recepção e corredores (off-white) · paredes internas do rooftop (cinza claro).",
         "Suvinil \"Vela\" / Coral \"Camurça Clara\" · rooftop: Suvinil \"Cromo\" / Coral \"Cinza Sublime\" (memorial de texturas).",
         "Briefing marca \"Outono Gelado\" (Coral) nos corredores: fechar a cor antes da compra. A VALIDAR.", RC1, (0.87, 0.35), 0.10),
        ("T16", "Rodapé", "Recepção, corredores e rooftop.",
         "Briefing e prancha 24 (ROD-01): poliestireno branco de sobrepor, quina viva · memorial de texturas: esmalte preto.",
         "As imagens mostram rodapé preto; briefing e executivo, branco. Definir antes da compra. A VALIDAR.", RC1, (0.70, 0.82), 0.10)]),
    ("Materiais · madeira e pedra (2/3)", [
        ("T17", "MDF Savana", "Sarrafos dos corredores, ripados e floreiras da recepção, ripados e prateleiras do rooftop.",
         "MDF Savana (amadeirado claro) com verniz fosco (memorial de texturas).",
         "Fixar após a pintura final; prever ponto das arandelas no eixo dos sarrafos.", RC1, (0.81, 0.45), 0.12),
        ("T18", "MDF Cinza Sagrado", "Armários do gourmet e do mini market do rooftop.",
         "MDF Cinza Sagrado (memorial de texturas).",
         "Memorial de interiores lista a marcenaria do rooftop como \"material a definir\". PENDENTE.", t(5), (0.12, 0.78), 0.15),
        ("T19", "Pedra moledo", "Painel do logo e parede do balanço, backsplash da copa, floreira e banco do rooftop.",
         "Pedra branca estilo moledo, irregular, junta seca (briefing e memorial de texturas).",
         "LOG Ind. 02: floreira e banco em alvenaria revestidos com a mesma pedra.", t(5), (0.12, 0.25), 0.15),
        ("T20", "Granito Pitaya", "Bancada da copa do rooftop.",
         "Granito Pitaya · cuba inox · torneira Gali bica alta (Docol).",
         "Bancada não instalada pela construtora (memorial 32.1): deixar pontos na cota do projeto.", t(5), (0.18, 0.515), 0.05)]),
    ("Materiais · rooftop externo e identidade (3/3)", [
        ("T21", "Deck do rooftop", "Deck aberto da cobertura e entorno da piscina.",
         "Porcelanato que imita madeira 120×20 (briefing) sobre piso elevado (prancha 44).",
         "Referência comercial antiderrapante não definida. PENDENTE.", t(1), (0.45, 0.88), 0.12),
        ("T22", "Revestimento da piscina", "Piscina aquecida da cobertura.",
         "Barlavento Blanc 14,5×14,5 (memorial, prancha 44) × briefing: revestimento esverdeado 20×20.",
         "Cor, formato e borda divergem entre memorial, prancha e briefing. A VALIDAR.", t(1), (0.20, 0.85), 0.12),
        ("T23", "Letreiros das áreas comuns", "Painel da recepção (sobre Velvet off-white) e painel em pedra do rooftop.",
         "\"Canasvieiras SPOT\" em letras pretas com LED no \"O\" (memorial de texturas).",
         "Briefing: ponto para logo iluminado com interruptor no logo; prever alimentação embutida antes do revestimento.",
         r(2), (0.55, 0.30), 0.18),
        ("T24", "Luminárias decorativas", "Corredores (arandelas) e gourmet do rooftop (pendentes).",
         "Arandela cilindro off-white com haste em Savana · pendente Bulan em rattan · spots de sobrepor pretos.",
         "Escopo Decor (memorial 32.1): deixar pontos elétricos no eixo dos sarrafos e sobre a mesa.", t(5), (0.70, 0.15), 0.20)])]
for title, cards in mats:
    S = B.dup(34)
    S.T(1060, "03 · INTERIORES DAS ÁREAS COMUNS · MAPA DE MATERIALIDADE")
    S.T(1061, title)
    S.T(1065, "Amostras extraídas das imagens de interiores Seazone. Referência visual, não substitui amostra física. "
              "Especificação: Memorial Descritivo, Memorial de Texturas, Memorial de Interiores e Briefing.")
    for ids, (code, ttl, onde, spec, obra, img, foc, hf) in zip(MC, cards):
        card(S, ids, code, ttl, onde, spec, obra, img, foc, hf)

# ================= 27 · ENCERRAMENTO =================
S = B.dup(38)
S.P(1168, F03, (0.6, 0.45))
S.T(1170, "CANASVIEIRAS SPOT")

B.finish(OUT, "CANASVIEIRAS SPOT")
print("ok", OUT)
